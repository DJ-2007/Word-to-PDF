#!/usr/bin/env python3
import os
import fitz # PyMuPDF
import docx
import openpyxl
from pptx import Presentation
from PIL import Image, ImageDraw

def create_samples(target_dir):
    os.makedirs(target_dir, exist_ok=True)
    
    # 1. TXT
    txt_path = os.path.join(target_dir, "sample.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("Convertly Universal File Converter\n=================================\n\nThis is a test plain text document containing multiple paragraphs.\nIt verifies formatting, typography, line wrapping, and UTF-8 encoding.\n\nThank you for using Convertly!\n")
    print(f"Created {txt_path}")
    
    # 2. CSV
    csv_path = os.path.join(target_dir, "sample.csv")
    with open(csv_path, "w", encoding="utf-8") as f:
        f.write("ID,Employee Name,Department,Salary,Status\n1,Alice Morgan,Engineering,115000,Active\n2,Bob Jenkins,Product,98000,Active\n3,Carol Wu,Design,92000,Active\n4,David Ross,Marketing,85000,Pending\n")
    print(f"Created {csv_path}")

    # 3. Markdown
    md_path = os.path.join(target_dir, "sample.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("# Convertly Engineering Whitepaper\n\n## Overview\nConvertly converts documents with **high fidelity** and zero data retention.\n\n- Real-time conversion\n- TLS 1.3 Transport\n- 24h Auto-deletion\n\n*Created for test verification.*\n")
    print(f"Created {md_path}")
    
    # 4. DOCX
    docx_path = os.path.join(target_dir, "sample.docx")
    d = docx.Document()
    d.add_heading('Quarterly Strategic Business Report', 0)
    p = d.add_paragraph('This document is generated to test Word to PDF conversion fidelity.')
    p.add_run(' Key metrics include 99.98% uptime and instantaneous conversion speeds.').bold = True
    d.add_heading('Executive Summary', level=1)
    d.add_paragraph('All core systems are operational and performing according to SLA targets.')
    d.save(docx_path)
    print(f"Created {docx_path}")
    
    # 5. XLSX
    xlsx_path = os.path.join(target_dir, "sample.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financials"
    ws.append(["Category", "Q1 Actual", "Q2 Actual", "Q3 Projected", "Q4 Target"])
    ws.append(["Software Infrastructure", 12500, 14200, 15000, 16000])
    ws.append(["Engineering Operations", 45000, 48000, 50000, 52000])
    ws.append(["Customer Support", 8000, 8500, 9000, 9500])
    wb.save(xlsx_path)
    print(f"Created {xlsx_path}")
    
    # 6. PPTX
    pptx_path = os.path.join(target_dir, "sample.pptx")
    prs = Presentation()
    slide_layout = prs.slide_layouts[0] # Title slide
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Convertly Presentation Deck"
    slide.placeholders[1].text = "Universal Cloud File Processing\nEngineered by Senior Full-Stack Team"
    # Slide 2
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Key Architectural Highlights"
    slide2.placeholders[1].text = "• Isolated temporary storage\n• Multi-adapter architecture\n• High performance Sharp and PyMuPDF engines"
    prs.save(pptx_path)
    print(f"Created {pptx_path}")
    
    # 7. Images (JPG, PNG)
    jpg1_path = os.path.join(target_dir, "sample1.jpg")
    img1 = Image.new('RGB', (400, 300), color=(79, 70, 229)) # Indigo
    d1 = ImageDraw.Draw(img1)
    d1.rectangle([(20, 20), (380, 280)], outline=(255, 255, 255), width=4)
    img1.save(jpg1_path, "JPEG", quality=90)
    print(f"Created {jpg1_path}")

    jpg2_path = os.path.join(target_dir, "sample2.jpg")
    img2 = Image.new('RGB', (400, 300), color=(16, 185, 129)) # Emerald
    d2 = ImageDraw.Draw(img2)
    d2.rectangle([(20, 20), (380, 280)], outline=(255, 255, 255), width=4)
    img2.save(jpg2_path, "JPEG", quality=90)
    print(f"Created {jpg2_path}")

    png1_path = os.path.join(target_dir, "sample1.png")
    img_png = Image.new('RGBA', (350, 350), color=(236, 72, 153, 200)) # Pink transparent
    img_png.save(png1_path, "PNG")
    print(f"Created {png1_path}")

    # 8. Multi-page PDF
    pdf_path = os.path.join(target_dir, "sample.pdf")
    doc = fitz.open()
    # Page 1
    page1 = doc.new_page(width=595, height=842) # A4
    page1.insert_text((50, 80), "CONVERTLY TEST PDF - PAGE 1", fontsize=20, color=(0.2, 0.2, 0.7))
    page1.insert_text((50, 120), "This document verifies multi-page extraction, rendering, splitting, and merging.", fontsize=11)
    page1.draw_rect(fitz.Rect(50, 150, 545, 300), color=(0.3, 0.5, 0.9), width=2)
    
    # Page 2
    page2 = doc.new_page(width=595, height=842)
    page2.insert_text((50, 80), "CONVERTLY TEST PDF - PAGE 2", fontsize=20, color=(0.1, 0.6, 0.3))
    page2.insert_text((50, 120), "Second page content for splitting and multi-page image export.", fontsize=11)
    
    doc.save(pdf_path)
    doc.close()
    print(f"Created {pdf_path}")

    # Second PDF for merging
    pdf2_path = os.path.join(target_dir, "sample2.pdf")
    doc2 = fitz.open()
    p = doc2.new_page(width=595, height=842)
    p.insert_text((50, 80), "APPENDIX A - ADDITIONAL ATTACHMENT", fontsize=18, color=(0.8, 0.2, 0.2))
    p.insert_text((50, 120), "This is the second PDF document to test PDF merging.", fontsize=11)
    doc2.save(pdf2_path)
    doc2.close()
    print(f"Created {pdf2_path}")

if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else ".temp_storage/samples"
    create_samples(target)
